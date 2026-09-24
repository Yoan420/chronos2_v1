from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import hashlib, json

ROOT=Path('C:/Users/BQ6757/chronos2_v1')
REF=Path('C:/Users/BQ6757/Downloads/Yoan_Kesraoui_CV.pdf')
EXPECTED='be82f5d11bbf8c2a432ee53be293399735230ec2822ba76234410f65ece43ab2'
assert hashlib.sha256(REF.read_bytes()).hexdigest()==EXPECTED
OUT=ROOT/'output/cv_citi_20260917/Yoan_Kesraoui_CV_Citi_Sales_Trading_2027.docx'
OUT.parent.mkdir(parents=True,exist_ok=True)
d=Document(); s=d.sections[0]
s.page_width=Pt(595.276); s.page_height=Pt(841.89)
s.left_margin=s.right_margin=Pt(43)
s.top_margin=Pt(36); s.bottom_margin=Pt(42)
s.header_distance=s.footer_distance=Pt(15)
DARK='15191E'; GRAY='4B525A'
for name in ('Normal','Title','Subtitle','Heading 1'):
    st=d.styles[name]
    st.font.name='Arial'; st.font.size=Pt(10); st.font.color.rgb=RGBColor.from_string(DARK)
    pf=st.paragraph_format
    pf.space_before=Pt(0); pf.space_after=Pt(0); pf.line_spacing=Pt(13.1)
    pf.widow_control=True
    rp=st.element.get_or_add_rPr()
    rf=rp.rFonts
    for a in ('ascii','hAnsi','cs','eastAsia'): rf.set(qn('w:'+a),'Arial')
    lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'en-GB'); rp.append(lang)
    for border in st.element.xpath('.//w:pBdr'): border.getparent().remove(border)
d.styles['Title'].font.size=Pt(23); d.styles['Title'].font.bold=True
d.styles['Title'].paragraph_format.line_spacing=Pt(27)
d.styles['Heading 1'].font.size=Pt(10.5); d.styles['Heading 1'].font.bold=True
d.styles['Heading 1'].paragraph_format.keep_with_next=True
d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
d.styles['Heading 1'].paragraph_format.space_after=Pt(4)

def para(text='',size=10,bold=False,italic=False,color=DARK,after=0,before=0,keep=False,align=None,line=13.1,style=None):
    p=d.add_paragraph(style=style)
    pf=p.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=Pt(line)
    pf.keep_with_next=keep; pf.keep_together=True
    if align is not None: p.alignment=align
    if text: run(p,text,size,bold,italic,color)
    return p

def run(p,text,size=10,bold=False,italic=False,color=DARK):
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(size); r.bold=bold; r.italic=italic
    r.font.color.rgb=RGBColor.from_string(color); return r

def link(p,label,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    for tag,attrs in [('rFonts',{'ascii':'Arial','hAnsi':'Arial'}),('sz',{'val':'18'}),('color',{'val':GRAY})]:
        el=OxmlElement('w:'+tag)
        for k,v in attrs.items(): el.set(qn('w:'+k),v)
        props.append(el)
    r.append(props); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); p._p.append(h)

def heading(text):
    return d.add_paragraph(text,'Heading 1')

def dated(title,date,size=10.1,before=0):
    p=para(before=before,keep=True,line=13.1)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(509.276),WD_TAB_ALIGNMENT.RIGHT)
    run(p,title,size=size,bold=True); run(p,'\t'+date,size=9.2,bold=True)
    return p

def secondary(text,after=2):
    return para(text,size=9.4,italic=True,color=GRAY,after=after,keep=True,line=12.8)

def bullet(label,text):
    p=para(after=3)
    p.paragraph_format.left_indent=Pt(10); p.paragraph_format.first_line_indent=Pt(-9)
    run(p,'• '); run(p,label+' ',bold=True); run(p,text)
    return p

p=para('YOAN KESRAOUI',size=23,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,line=27,style='Title',after=1)
para('SALES AND TRADING  |  SHORT-TERM POWER  |  RISK ANALYTICS',size=10.3,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,line=13.5,after=5)
p=para(align=WD_ALIGN_PARAGRAPH.CENTER,line=12.4,after=0)
run(p,'Paris, France  |  +33 7 81 83 13 88  |  ',size=9.2,color=GRAY)
link(p,'kesraoui.yoan@gmail.com','mailto:kesraoui.yoan@gmail.com')
run(p,'  |  ',size=9.2,color=GRAY)
link(p,'LinkedIn','https://www.linkedin.com/in/yoan-kesraoui-20b983304/')

heading('PROFILE')
para("ML/DL researcher developing forecasts and analytical tools for ENGIE's short-term power trading desk, with experience in front-office market risk, pricing and market microstructure. Seeking Citi's 2027 Markets Sales and Trading off-cycle internship in Paris.")
para('Available from January 2027 for six months | University internship agreement confirmed.',size=9.5,bold=True,before=2,line=12.6)

heading('EDUCATION')
dated('Université Paris 1 Panthéon-Sorbonne - M.Sc. Data Science','2025-2026',size=9.9)
secondary('Machine learning, deep learning, statistics, optimisation and time series',after=4)
dated('Université Paris Cité - M.Sc. Applied Economics - Econometrics','2023-2025',size=9.9)
secondary('Econometrics, statistical inference, financial modelling and panel data',after=4)
dated('Université Paris Cité - B.Sc. Economics & Mathematics','2020-2023',size=9.9)

heading('PROFESSIONAL EXPERIENCE')
dated('ENGIE Global Markets - ML/DL Research Intern - Power Market','Jun 2026-Present')
secondary('Global Market Analysis, Paris | Research and tools for the short-term power trading desk')
bullet('Desk analytics:',"Developed NYX for day-ahead power-price forecasting, combining Chronos-2, CatBoost residual models and adaptive Kalman filtering.")
bullet('Forecast evaluation:', 'Reduced pooled MAE by 8.27% versus Chronos-2 alone in a retrospective 350-day evaluation across four European markets; assessed price spikes, negative prices and forecast uncertainty.')
bullet('Power fundamentals:', 'Analysed residual load, cross-border flows, nuclear availability and supply stacks to explain day-ahead price dynamics under European market coupling.')
bullet('Research communication:', 'Built rolling validation and calibration workflows; documented NYX methods, results and limitations in an unpublished English research manuscript.')

dated('Société Générale - Quant Market Risk Intern','Sep 2025-Mar 2026',before=5)
secondary('Front Office Market Risk, Paris')
bullet('Risk modelling:', 'Built GARCH, EGARCH and VAR models; supported Greeks, VaR and stress testing.')
bullet('Data and controls:', 'Built Python/SQL/API pipelines for large risk datasets; applied clustering and KNN to anomaly detection and risk monitoring.')

dated('Autorité des marchés financiers (AMF) - Data Scientist / Economist','Jan-Jun 2025',before=5)
secondary('Financial Stability & Risk Studies, Paris')
bullet('Market microstructure:', 'Analysed millions of MiFID II transactions using PySpark, Python, R and panel econometrics to study investor behaviour and market dynamics.')
bullet('Research output:', 'Engineered investor-level turnover, return and trading-intensity indicators; contributed empirical analysis to the AMF Market and Risk Mapping 2025.')

dated('Generali - Quantitative Analyst Intern','Apr-Sep 2024',before=5)
secondary('ALM & Strategic Asset Allocation, Paris')
bullet('Pricing:', 'Built XGBoost surrogate models for computationally intensive full-revaluation pricing; investigated non-linear relationships between risk factors.')
bullet('Rates:', 'Calibrated Vasicek/CIR models; automated risk and ALM workflows using Python, SQL and Power BI.')

heading('MARKETS AND TECHNICAL SKILLS')
for label,text in [
    ('Markets and risk:', 'Power, oil, derivatives, rates, pricing, Greeks, VaR/stress testing, ALM, market microstructure'),
    ('Programming and data:', 'Python, SQL, R, PySpark/Spark, Git, pandas, NumPy, statsmodels, Power BI'),
    ('ML/DL:', 'PyTorch, TensorFlow, scikit-learn, CatBoost, XGBoost'),
    ('Methods:', 'Time-series forecasting, econometrics, Monte Carlo, optimisation, backtesting, anomaly detection'),
    ('Languages:', 'French (native), English (C1), Spanish (beginner)')]:
    p=para(size=9.5,after=2,line=12.6)
    run(p,label+' ',size=9.5,bold=True); run(p,text,size=9.5)

d.core_properties.title='Yoan Kesraoui - Citi Markets Sales and Trading Internship 2027'
d.core_properties.subject='Short-term power trading desk research and market risk'
d.core_properties.author='Yoan Kesraoui'
d.core_properties.keywords='Citi, Sales and Trading, power, NYX, Python, market risk'
d.core_properties.comments=''
d.save(OUT)
assert hashlib.sha256(REF.read_bytes()).hexdigest()==EXPECTED
print(OUT)
print('Words:',len(' '.join(p.text for p in d.paragraphs).split()))
