from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'cv_these_75968'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.25), Cm(1.25)
sec.left_margin, sec.right_margin = Cm(1.55), Cm(1.55)
sec.header_distance = sec.footer_distance = Cm(.4)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(10.5)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.line_spacing = 1.03
    s.paragraph_format.widow_control = True
    s.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    s.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'fr-FR')
    s.element.get_or_add_rPr().append(lang)

normal = doc.styles['Normal']
normal.paragraph_format.space_after = Pt(2)
title = doc.styles['Title']
title.font.size = Pt(25)
title.font.bold = True
title.paragraph_format.space_after = Pt(3)
h1 = doc.styles['Heading 1']
h1.font.size = Pt(11)
h1.font.bold = True
h1.font.all_caps = False
h1.paragraph_format.space_before = Pt(9)
h1.paragraph_format.space_after = Pt(4)
h1.paragraph_format.keep_with_next = True

def para(text='', *, size=None, bold=False, italic=False, after=2, before=0, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    return p

def hyperlink(p, text, url):
    el = OxmlElement('w:hyperlink')
    el.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    for tag, attrs in [('w:rFonts', {'w:ascii':'Calibri','w:hAnsi':'Calibri'}),
                       ('w:color', {'w:val':'333333'}), ('w:sz', {'w:val':'19'})]:
        e=OxmlElement(tag)
        for k,v in attrs.items(): e.set(qn(k),v)
        props.append(e)
    run.append(props)
    t=OxmlElement('w:t'); t.text=text; run.append(t); el.append(run); p._p.append(el)

def dated(text, date, before=0):
    p = para(after=1, before=before, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.9), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(text).bold=True
    r=p.add_run('\t'+date); r.font.size=Pt(9.5)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph()
    pf=p.paragraph_format
    pf.left_indent=Cm(.28); pf.first_line_indent=Cm(-.28)
    pf.space_after=Pt(2); pf.keep_together=True
    p.add_run('• ')
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True
        p.add_run(text[len(bold_prefix):])
    else: p.add_run(text)
    return p

p=doc.add_paragraph('YOAN KESRAOUI', 'Title')
para('Candidature doctorale', size=10.5, after=1)
para('Modèles de fondation tabulaires pour les données hétérogènes', size=12, bold=True, after=2)
para('Équipe SODA  ·  Inria Saclay  ·  Institut Polytechnique de Paris', size=9.5, after=4)
p=para(size=9.5, after=1)
r=p.add_run('Paris, France  ·  +33 7 81 83 13 88  ·  '); r.font.size=Pt(9.5)
hyperlink(p,'kesraoui.yoan@gmail.com','mailto:kesraoui.yoan@gmail.com')
p=para(after=7)
hyperlink(p,'linkedin.com/in/yoan-kesraoui-20b983304','https://www.linkedin.com/in/yoan-kesraoui-20b983304/')

para('Formation en data science, économétrie et mathématiques. Travaux sur les mécanismes des Transformers, utilisation de Chronos-2 et évaluation de modèles sur données réelles. Projet de recherche : modèles de fondation tabulaires et apprentissage sur données hétérogènes.', after=0)

doc.add_paragraph('Formation', 'Heading 1')
dated('Université Paris 1 Panthéon-Sorbonne  |  M.Sc. Data Science', '2025 - 2026')
para('Machine learning, deep learning, statistiques, optimisation et séries temporelles.', size=10, after=4)
dated('Université Paris Cité  |  M.Sc. Applied Economics - Econometrics', '2023 - 2025')
para('Économétrie, inférence statistique, modélisation financière et données de panel.', size=10, after=4)
dated('Université Paris Cité  |  B.Sc. Economics & Mathematics', '2020 - 2023')

doc.add_paragraph('Mémoire de fin d’études', 'Heading 1')
para('Transformers et modèles de fondation pour les séries temporelles', bold=True, after=2)
bullet('Revue de littérature : attention, tokenisation, masquage et préentraînement ; analyse des architectures temporelles et des modèles de fondation, dont Chronos-2, TimesFM et Moirai.')
bullet('Analyse critique du transfert, des fuites de données et des prévisions probabilistes ; proposition de protocoles expérimentaux sur la généralisation, les covariables et l’adaptation.')

doc.add_paragraph('Expériences de recherche et de modélisation', 'Heading 1')
dated('ENGIE Global Markets  |  Stagiaire en recherche quantitative', 'Juin 2026 - présent')
para('Global Market Analysis, Paris', size=9.5, italic=True, after=2)
bullet('Prévision des prix de l’électricité française à J+1 en combinant Chronos-2, LEAR, CatBoost et ensembles hors pli ; MAE de 12,82 €/MWh sur un jeu de test final réservé de 365 jours.')
bullet('Pipelines respectant l’information disponible à chaque date, validation glissante et contrôles des fuites de données ; analyse des biais, des événements extrêmes et des changements de régime.')

dated('Société Générale  |  Stagiaire en risque de marché quantitatif', 'Sept. 2025 - mars 2026', before=5)
para('Front Office Market Risk, Paris', size=9.5, italic=True, after=2)
bullet('Modèles GARCH, EGARCH et VAR ; pipelines Python/SQL/API sur de grands jeux de données ; clustering et k plus proches voisins pour la détection d’anomalies.')

dated('Autorité des marchés financiers  |  Data scientist et économiste', 'Janv. - juin 2025', before=5)
para('Études sur la stabilité financière et les risques, Paris', size=9.5, italic=True, after=2)
bullet('Analyse de millions de transactions MiFID II avec PySpark, Python, R et économétrie de panel ; construction d’indicateurs individuels et étude des comportements des investisseurs.')
bullet('Contribution aux analyses empiriques de la Cartographie des marchés et des risques 2025 de l’AMF.')

dated('Generali  |  Stagiaire analyste quantitatif', 'Avr. - sept. 2024', before=5)
para('Gestion actif-passif et allocation stratégique, Paris', size=9.5, italic=True, after=2)
bullet('Développement de modèles de substitution XGBoost pour des valorisations coûteuses en calcul ; calibration de modèles de taux Vasicek/CIR et analyse de relations non linéaires entre facteurs de risque.')

doc.add_paragraph('Compétences scientifiques et techniques', 'Heading 1')
for label,text in [
    ('Apprentissage automatique  ', 'PyTorch, TensorFlow, scikit-learn, Chronos-2, CatBoost, XGBoost.'),
    ('Programmation et données  ', 'Python, SQL, R, PySpark/Spark, pandas, NumPy, statsmodels, Git.'),
    ('Méthodes  ', 'Modélisation probabiliste, économétrie, optimisation, Monte-Carlo, validation temporelle.'),
    ('Langues  ', 'Français natif, anglais C1, espagnol débutant.'),
]:
    p=para(after=2)
    p.add_run(label).bold=True
    p.add_run(text)

doc.core_properties.title='Yoan Kesraoui - Candidature doctorale en modèles de fondation tabulaires'
doc.core_properties.subject='Candidature SODA Inria Saclay - Offre ADUM 75968'
doc.core_properties.author='Yoan Kesraoui'
doc.core_properties.keywords='data science, modèles de fondation, données tabulaires, Python, PyTorch'
doc.core_properties.comments=''
for paragraph in doc.paragraphs:
    if paragraph.style.name == 'Heading 1':
        for run in paragraph.runs:
            run.text = run.text.upper()
doc.save(OUT / 'Yoan_Kesraoui_CV_These_SODA.docx')
print(OUT / 'Yoan_Kesraoui_CV_These_SODA.docx')
print('Words:', len(' '.join(p.text for p in doc.paragraphs).split()))
