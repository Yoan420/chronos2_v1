from pathlib import Path
import sys, re, json
from copy import deepcopy
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE/'deps_readable'))
from latex2mathml.converter import convert
OUT=BASE.parents[1]/'output'/'cv_these_75968'
transform=etree.XSLT(etree.parse(r'C:\Program Files (x86)\Microsoft Office\Root\Office16\MML2OMML.XSL'))
doc=Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
section=doc.sections[0]
section.page_width,section.page_height=Cm(21),Cm(29.7)
section.left_margin=section.right_margin=Cm(2.2)
section.top_margin=Cm(2)
section.bottom_margin=Cm(1.9)
section.header_distance=section.footer_distance=Cm(.8)
section.different_first_page_header_footer=True

for s in doc.styles:
    if s.type==1:
        s.font.name='Cambria'
        s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.widow_control=True
        rp=s.element.get_or_add_rPr()
        rf=rp.find(qn('w:rFonts'))
        if rf is None: rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        rf.set(qn('w:ascii'),'Cambria');rf.set(qn('w:hAnsi'),'Cambria')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rf.attrib.pop(qn('w:'+key),None)
        lang=OxmlElement('w:lang');lang.set(qn('w:val'),'fr-FR');rp.append(lang)

normal=doc.styles['Normal']
normal.font.size=Pt(11.5)
normal.paragraph_format.line_spacing=1.12
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name,size,before,after in [('Heading 1',18,0,14),('Heading 2',13,12,7),('Heading 3',11.5,9,5)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
    pf=s.paragraph_format;pf.space_before=Pt(before);pf.space_after=Pt(after)
    pf.keep_with_next=True;pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Heading 1'].paragraph_format.page_break_before=False
doc.styles['Heading 1'].paragraph_format.space_before=Pt(20)
front=doc.styles.add_style('Front heading',WD_STYLE_TYPE.PARAGRAPH)
front.base_style=normal;front.font.size=Pt(18);front.font.bold=True
front.paragraph_format.page_break_before=True
front.paragraph_format.space_after=Pt(14)
front.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
front.paragraph_format.keep_with_next=True
doc.styles['Title'].font.size=Pt(27)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.line_spacing=1.05
doc.styles['Title'].paragraph_format.space_after=Pt(18)
doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
caption=doc.styles['Caption'];caption.font.size=Pt(9.5);caption.font.italic=True
caption.paragraph_format.space_after=Pt(8)
caption.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
for name in ['TOC 1','TOC 2']:
    if name in doc.styles:
        s=doc.styles[name];s.font.size=Pt(10.5)
        s.paragraph_format.space_after=Pt(3)
        s.paragraph_format.line_spacing=1

foot=section.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
foot._p.append(field)

def add_link(p,label,url):
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'1F4E79');props.append(color)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);p._p.append(link)

eq_log=[]
def add_math(p,latex,display=False):
    latex=latex.strip()
    tag=re.search(r'\\tag\{([^}]+)\}',latex)
    number=tag.group(1) if tag else ''
    latex=re.sub(r'\\tag\{[^}]+\}','',latex)
    if r'\qquad' in latex:
        for i,segment in enumerate(latex.split(r'\qquad')):
            if i:p.add_run('   ')
            if segment.strip()==r'\text{ou}':p.add_run('ou')
            else:add_math(p,segment,display)
        if number:p.add_run('    ('+number+')')
        eq_log.append({'latex':latex,'label':number})
        return
    mathml=convert(latex,display='block' if display else 'inline')
    omml=transform(etree.fromstring(mathml.encode())).getroot()
    # Set an explicit math font and retain native, editable Office equations.
    for run in omml.xpath('.//m:r',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}):
        value=''.join(run.itertext())
        if value in ['Attention','softmax','MHA','Concat','FFN','ReLU','LN','MAE','RMSE','MASE','CRPS','model','pos','ff','ou']:
            mathprops=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');mathprops.append(sty)
            run.insert(0,mathprops)
        props=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');props.append(fonts)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'23');props.append(sz);run.insert(0,props)
    p._p.append(deepcopy(omml))
    if number: p.add_run('    ('+number+')')
    eq_log.append({'latex':latex,'label':number})

pattern=re.compile(r'(\*\*.+?\*\*|\[[^\]]+\]\(https?://[^)]+\)|\$[^$]+\$)')
def inline(p,text):
    text=text.replace('μ_c', '$\\mu_c$').replace('xτ', '$x_\\tau$').replace('vθ', '$v_\\theta$').replace('Bℓ', '$B_\\ell$')
    for part in pattern.split(text):
        if part.startswith('**') and part.endswith('**'):p.add_run(part[2:-2]).bold=True
        elif part.startswith('[') and '](' in part:
            m=re.match(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
            add_link(p,m.group(1),m.group(2))
        elif part.startswith('$') and part.endswith('$'):add_math(p,part[1:-1])
        else:
            for fragment in re.split(r'(\b[A-Za-z]_(?:model|ff|pos|[A-Za-z0-9])\b)',part):
                if re.fullmatch(r'[A-Za-z]_(?:model|ff|pos|[A-Za-z0-9])',fragment):
                    a,b=fragment.split('_',1)
                    if len(b)>1:b=r'\mathrm{'+b+'}'
                    add_math(p,a+'_{'+b+'}')
                else:p.add_run(fragment)

def body(text,style=None):
    p=doc.add_paragraph(style=style)
    inline(p,text)
    return p

def render_md(md,bibliography=False):
    md=re.sub(r'(?m)^(#{1,3} .+)$',r'\n\1\n',md)
    blocks=re.split(r'\n\s*\n',md.strip())
    for b in blocks:
        b=b.strip()
        if not b:continue
        if b.startswith('$$'):
            latex=b[2:]
            pos=latex.rfind('$$')
            trailing=latex[pos+2:].strip() if pos>=0 else ''
            latex=latex[:pos] if pos>=0 else latex
            if trailing and not '\\tag{' in latex:
                label=re.search(r'\((\d+[a-z]?)\)',trailing)
                if label:latex+='\\tag{'+label.group(1)+'}'
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(6)
            p.paragraph_format.keep_together=True
            add_math(p,latex,True)
        elif b.startswith('[[FIGURE_'):
            num=re.search(r'FIGURE_(\d+)',b).group(1)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            run=p.add_run();run.add_picture(str(BASE/f'figure_{num}.png'),width=Cm(16.4))
            props=run._r.xpath('.//wp:docPr')
            if props:props[0].set('descr',f'Figure {num} du mémoire illustrant les mécanismes des Transformers')
            rest=b[b.index(']]')+2:].strip()
            if rest:body(rest,'Caption')
        elif b.startswith('# '):
            title=b[2:].replace('\n',' ')
            p=doc.add_paragraph(title,'Heading 1')
            if title=='1 Introduction':p.paragraph_format.page_break_before=True
        elif b.startswith('## '):doc.add_paragraph(b[3:].replace('\n',' '),'Heading 2')
        elif b.startswith('### '):doc.add_paragraph(b[4:].replace('\n',' '),'Heading 3')
        elif b.startswith('Figure '):body(b.replace('\n',' '),'Caption')
        else:
            p=body(b.replace('\n',' '))
            if bibliography:
                p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_together=True
                p.paragraph_format.space_after=Pt(8)
                for r in p.runs:r.font.size=Pt(10.5)

# Cover, preserving the subject and administrative information supplied by the author.
p=doc.add_paragraph('SORBONNE DATA ANALYTICS');p.paragraph_format.space_after=Pt(48)
p.runs[0].font.size=Pt(12);p.runs[0].bold=True
p=doc.add_paragraph('Mémoire de fin d’études');p.paragraph_format.space_after=Pt(16)
p.runs[0].font.size=Pt(15)
doc.add_paragraph('Transformers et modèles de fondation pour les séries temporelles','Title')
p=doc.add_paragraph('Représentations, préentraînement et conditions du transfert')
p.runs[0].font.size=Pt(15);p.paragraph_format.space_after=Pt(38)
p=doc.add_paragraph('YOAN KESRAOUI');p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
doc.add_paragraph('Année universitaire 2025 - 2026')
doc.add_paragraph('Numéro étudiant 12517303')
p=doc.add_paragraph('Septembre 2026');p.paragraph_format.space_before=Pt(32)

revision=(BASE/'revision_recherche.md').read_text(encoding='utf-8')
summary=revision.split('# Résumé\n',1)[1].split('# Abstract\n',1)[0].strip()
abstract=revision.split('# Abstract\n',1)[1].split('# 1 Introduction\n',1)[0].strip()
intro='# 1 Introduction\n\n'+revision.split('# 1 Introduction\n',1)[1].split('# 6 Pistes',1)[0].strip()
chapter6='# 6 Pistes '+revision.split('# 6 Pistes',1)[1].split('# 7 Conclusion\n',1)[0].strip()
conclusion='# 7 Conclusion\n\n'+revision.split('# 7 Conclusion\n',1)[1].split('# Références à ajouter',1)[0].strip()

doc.add_paragraph('Résumé','Front heading');render_md(summary)
p=doc.add_paragraph('Abstract');p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
p.paragraph_format.space_before=Pt(16);p.paragraph_format.keep_with_next=True
render_md(abstract)

doc.add_paragraph('Sommaire','Front heading')
p=doc.add_paragraph()
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
p._p.append(field)

render_md(intro)
render_md((BASE/'memoire_ch2_5.md').read_text(encoding='utf-8'))
render_md(chapter6)
render_md(conclusion)
bib=(BASE/'bibliographie.md').read_text(encoding='utf-8')
bibblocks=re.split(r'\n\s*\n',bib.strip())
p=doc.add_paragraph('Bibliographie','Heading 1');p.paragraph_format.page_break_before=True
body('Les références sont classées par nom du premier auteur. Les statuts de publication et les versions précisent la portée des sources. Les sources en ligne du corpus initial ont été consultées les 9 et 10 septembre 2026.')
render_md('\n\n'.join(bibblocks[2:]),True)

doc.core_properties.title='Transformers et modèles de fondation pour les séries temporelles'
doc.core_properties.author='Yoan Kesraoui'
doc.core_properties.subject='Mémoire de fin d’études Sorbonne Data Analytics'
doc.core_properties.keywords='Transformers, séries temporelles, modèles de fondation, prévision probabiliste'
doc.core_properties.comments=''
for p in doc.paragraphs:
    # Never inherit a decorative paragraph border from the base document.
    for border in p._p.xpath('.//w:pBdr'):border.getparent().remove(border)
path=OUT/'Memoire_Yoan_Kesraoui_Revise.docx'
doc.save(path)
(BASE/'equations_qa.json').write_text(json.dumps(eq_log,ensure_ascii=False,indent=2),encoding='utf-8')
print(path)
print('Equations',len(eq_log),'Paragraphs',len(doc.paragraphs))
