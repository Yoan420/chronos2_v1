from pathlib import Path
from pypdf import PdfReader
from docx import Document
import re, json
base=Path(__file__).resolve().parent
for stem in ['Yoan_Kesraoui_CV_These_SODA','Memoire_Yoan_Kesraoui_Revise']:
    root=base.parents[1]/'output'/'cv_these_75968'
    d=Document(root/(stem+'.docx'))
    p=PdfReader(root/(stem+'.pdf'))
    full='\n'.join(x.extract_text() or '' for x in p.pages)
    report={'file':stem,'pages':len(p.pages),'words_per_page':[len((x.extract_text() or '').split()) for x in p.pages],
            'formulas':len(d.element.xpath('.//m:oMath')),'figures':len(d.inline_shapes),
            'headings':[(x.style.name,x.text) for x in d.paragraphs if x.style.name.startswith('Heading')]}
    assert not re.search(r'(^|\s)##',full), 'Markdown heading left in PDF'
    if stem.startswith('Memoire'):
        assert not any(s in full for s in ['TabPFN','TabICL','H5']), 'Removed research direction remains'
        assert len(d.inline_shapes)==3
        assert len([r for r in d.part.rels.values() if 'hyperlink' in r.reltype])==41
        assert any(x.text=='1.1 Sujet et problématique' for x in d.paragraphs)
        assert any(x.text=='7 Conclusion' for x in d.paragraphs)
        assert 'aucune des expériences proposées' in full
    else:
        assert len(p.pages)==1
        for token in ['kesraoui.yoan@gmail.com','+33 7 81 83 13 88','12,82','365','Revue de littérature']:
            assert token in full,token
    (base/(stem+'_qa.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(stem,'pages',len(p.pages),'words per page',report['words_per_page'])
