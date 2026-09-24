from copy import deepcopy
from pathlib import Path
import zipfile
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path('C:/Users/BQ6757/chronos2_v1')
OUT = ROOT/'output/candidature_inria_sierra_20260917'
OUT.mkdir(parents=True, exist_ok=True)
SOURCE = ROOT/'output/cv_soda_memoire_sg_20260917/Yoan_Kesraoui_CV_These_SODA_NYX.docx'
CV = OUT/'Yoan_Kesraoui_CV_Inria_SIERRA_Gradient_Stochastique.docx'
LETTER = OUT/'Yoan_Kesraoui_Lettre_Motivation_Inria_SIERRA.docx'

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
NS = {'w': W[1:-1]}

with zipfile.ZipFile(SOURCE) as z:
    files = {name: z.read(name) for name in z.namelist()}
root = E.fromstring(files['word/document.xml'])
body = root.find('w:body', NS)
ps = body.findall('w:p', NS)

def text(p):
    return ''.join(p.xpath('.//w:t/text()', namespaces=NS))

def find(prefix):
    candidates = [p for p in ps if text(p).startswith(prefix)]
    assert len(candidates) == 1, (prefix, len(candidates))
    return candidates[0]

def rewrite(p, new):
    run = next((r for r in p.findall('w:r', NS) if r.find('w:t', NS) is not None), None)
    prop = deepcopy(run.find('w:rPr', NS)) if run is not None and run.find('w:rPr', NS) is not None else None
    for child in list(p):
        if child.tag != W+'pPr': p.remove(child)
    r = E.SubElement(p, W+'r')
    if prop is not None: r.append(prop)
    t = E.SubElement(r, W+'t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = new

rewrite(find('Candidature '), 'Candidature doctorale')
rewrite(find('Modèles de fondation tabulaires'), 'Estimation de la performance du gradient stochastique')
rewrite(find('Équipe SODA'), 'Équipe SIERRA · Centre Inria de Paris · Direction : Francis Bach')

sg_title = find('Mémoire | Société Générale')
sg_detail = body[body.index(sg_title)+1]
assert 'Sridi et Bilokon' in text(sg_detail)
rewrite(sg_title, 'Mémoire | Société Générale — Apprentissage différentiel et calibration de Heston')
rewrite(sg_detail, '• Étude et implémentation de réseaux denses entraînés sur les prix de puts européens et leurs sensibilités ; calibration du modèle de Heston (d’après Sridi et Bilokon, 2023).')
nyx_title = find('NYX | Travaux réalisés')
body.remove(sg_title)
body.remove(sg_detail)
index = body.index(nyx_title)
body.insert(index, sg_title)
body.insert(index+1, sg_detail)

methods = find('Méthodes  ')
method_runs = methods.findall('w:r', NS)
assert len(method_runs) == 3 and text(methods).startswith('Méthodes  Transformers')
method_runs[1].find('w:t', NS).text = 'Optimisation, économétrie, évaluation empirique, validation temporelle, Monte-Carlo.'
methods.remove(method_runs[2])
rewrite(find('• Modèles GARCH'), '• Modèles GARCH, EGARCH et VAR ; pipelines Python/SQL/API sur de grands jeux de données ; clustering et distances aux plus proches voisins pour la détection d’anomalies.')

files['word/document.xml'] = E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
with zipfile.ZipFile(CV, 'w', zipfile.ZIP_DEFLATED) as z:
    for name, content in files.items(): z.writestr(name, content)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.95)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.25)
sec.right_margin = Cm(2.25)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.8)
normal.font.color.rgb = RGBColor(28, 32, 37)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.1

def paragraph(text='', bold=False, size=None, after=None, before=None, align=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if before is not None: p.paragraph_format.space_before = Pt(before)
    run = p.add_run(text)
    if bold: run.bold = True
    if size: run.font.size = Pt(size)
    return p

paragraph('YOAN KESRAOUI', bold=True, size=15, after=2)
paragraph('Paris, France  ·  +33 7 81 83 13 88  ·  kesraoui.yoan@gmail.com', size=9.5, after=1)
paragraph('linkedin.com/in/yoan-kesraoui-20b983304', size=9.5, after=12)
paragraph('Paris, le 17 septembre 2026', after=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
paragraph('À l’attention de Monsieur Francis Bach\nÉquipe SIERRA · Centre Inria de Paris', after=11)
paragraph('Objet : Candidature au doctorat « Estimation de la performance du gradient stochastique » (offre 2026-10295)', bold=True, after=12)
paragraph('Monsieur Bach,', after=8)

paragraph('Je vous adresse ma candidature à la thèse consacrée à l’estimation de la performance du gradient stochastique au sein de l’équipe SIERRA. Ce sujet correspond à mon souhait d’approfondir les fondements statistiques et mathématiques des méthodes d’apprentissage que j’ai jusqu’ici étudiées et mises en œuvre dans des contextes appliqués.', after=9)

paragraph('Ma formation associe un M.Sc. en Data Science à l’Université Paris 1 Panthéon-Sorbonne et un M.Sc. en économie appliquée et économétrie à l’Université Paris Cité. Elle m’a donné des bases en statistique, économétrie, optimisation et apprentissage automatique. Chez Société Générale, j’ai étudié et implémenté une approche d’apprentissage différentiel pour le pricing d’options sous Heston : le réseau apprend les prix et leurs sensibilités, qui servent ensuite à la calibration du modèle. Ce travail m’a amené à examiner la précision des gradients et la stabilité d’une méthode d’optimisation sur un problème concret.', after=9)

paragraph('Dans mon stage actuel chez ENGIE Global Markets, je conçois NYX, une méthode de prévision des prix de l’électricité associant un modèle de fondation, une correction des résidus et un filtrage adaptatif. Son évaluation rétrospective sur quatre marchés et 350 jours, avec une validation chronologique, m’a sensibilisé à la manière dont un protocole expérimental permet de juger la performance réelle d’une méthode. J’aimerais désormais compléter cette pratique par un travail de recherche qui analyse plus directement les mécanismes et les garanties des algorithmes d’apprentissage.', after=9)

paragraph('L’articulation entre optimisation et apprentissage statistique dans les travaux de SIERRA me paraît particulièrement stimulante pour cette transition. Je serais heureux d’échanger avec vous sur le sujet et sur la manière dont mon parcours pourrait y contribuer. Je vous remercie de l’attention portée à ma candidature.', after=12)
paragraph('Je vous prie d’agréer, Monsieur Bach, l’expression de mes salutations distinguées.', after=17)
paragraph('Yoan Kesraoui', bold=True)
doc.save(LETTER)

assert CV.is_file() and LETTER.is_file()
print(CV)
print(LETTER)
