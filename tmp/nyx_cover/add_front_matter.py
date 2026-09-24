from pathlib import Path
from io import BytesIO
from copy import deepcopy
import zipfile, hashlib, json
from lxml import etree as E
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('C:/Users/BQ6757/chronos2_v1')
TMP = ROOT / 'tmp/nyx_cover'
SRC = ROOT / 'output/nyx_paper_20250831_20260831/NYX_Research_Manuscript.docx'
OUT = ROOT / 'output/nyx_paper_with_cover/NYX_Research_Manuscript.docx'
EXPECTED = 'ebbaff8d0bf7701a95159ec954261c7505569ade482210d95e9220dab1e5db8f'
assert hashlib.sha256(SRC.read_bytes()).hexdigest() == EXPECTED, 'Source changed; rebase before editing.'
NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'

def el(tag, **attrs):
    a = OxmlElement('w:'+tag)
    for k,v in attrs.items(): a.set(qn('w:'+k), str(v))
    return a

def fmt(p, size=12, bold=False, italic=False, before=0, after=0, align=None):
    f=p.paragraph_format
    f.space_before=Pt(before); f.space_after=Pt(after)
    f.line_spacing=1.0
    f.keep_with_next=False; f.keep_together=True
    if align is not None: p.alignment=align
    pp=p._p.get_or_add_pPr()
    pp.append(el('outlineLvl',val=9))
    for r in p.runs:
        r.font.name='Times New Roman'; r.font.size=Pt(size)
        r.font.bold=bold; r.font.italic=italic
        r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    return p

def para(txt='', **kw): return fmt(d.add_paragraph(txt), **kw)

d=Document()
s=d.sections[0]
s.page_width=Pt(612); s.page_height=Pt(792)
s.left_margin=Pt(59.4); s.right_margin=Pt(52.6)
s.top_margin=Pt(28.9); s.bottom_margin=Pt(45)
s.header_distance=Pt(0); s.footer_distance=Pt(20)

# Native images in one top band. Preserve source image bytes and aspect ratios.
p=para(after=32)
p.paragraph_format.tab_stops.add_tab_stop(Pt(500), WD_TAB_ALIGNMENT.RIGHT)
p.add_run().add_picture(str(TMP/'sorbonne_logo.png'), width=Pt(265))
p.add_run('\t')
p.add_run().add_picture(str(TMP/'engie_logo.png'), width=Pt(166))
fmt(p, after=32)

p=para('NAME\t: KESRAOUI Yoan\tAcademic year 2025–2026',bold=True,after=14)
p.paragraph_format.tab_stops.add_tab_stop(Pt(119))
p.paragraph_format.tab_stops.add_tab_stop(Pt(277))
p=para('Student number\t: 12517303',bold=True,after=37)
p.paragraph_format.tab_stops.add_tab_stop(Pt(119))
para('SORBONNE DATA ANALYTICS',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=37)
para('RESEARCH MANUSCRIPT',size=14,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=22)

t=d.add_table(rows=1, cols=1); t.autofit=False
t.columns[0].width=Pt(500); t.cell(0,0).width=Pt(500)
pr=t._tbl.tblPr
pr.find(qn('w:tblW')).set(qn('w:w'),'10000')
pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
pr.append(el('tblInd',w=150,type='dxa'))
borders=el('tblBorders')
for edge in ('top','left','bottom','right','insideH','insideV'):
    borders.append(el(edge,val='single',sz=4,color='000000'))
pr.append(borders)
cell=t.cell(0,0); cp=cell._tc.get_or_add_tcPr()
cp.append(el('shd',val='clear',fill='D9D9D9'))
marg=el('tcMar')
for k,v in [('top',120),('bottom',140),('left',150),('right',150)]: marg.append(el(k,w=v,type='dxa'))
cp.append(marg)
p=cell.paragraphs[0]; p.style='Title'
p.add_run('NYX for hourly day ahead electricity price forecasting across CWE')
fmt(p,size=14,bold=True,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
p=cell.add_paragraph('Foundation models, residual learning and adaptive filtering')
fmt(p,size=14,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
para('Study period 31 August 2025 to 31 August 2026',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,before=19)

def section_end(margins):
    p=para('',size=1)
    p.paragraph_format.line_spacing=Pt(1)
    sp=deepcopy(s._sectPr)
    for a in list(sp):
        if E.QName(a).localname in ('headerReference','footerReference','pgNumType','titlePg'): sp.remove(a)
    typ=sp.find(qn('w:type'))
    if typ is None: typ=el('type',val='nextPage'); sp.insert(0,typ)
    else: typ.set(qn('w:val'),'nextPage')
    mg=sp.find(qn('w:pgMar'))
    for k,v in margins.items(): mg.set(qn('w:'+k),str(round(v*20)))
    p._p.get_or_add_pPr().append(sp)

section_end({'top':28.9,'bottom':45,'left':59.4,'right':52.6})

p=para(align=WD_ALIGN_PARAGRAPH.RIGHT,after=8)
p.add_run().add_picture(str(TMP/'engie_logo.png'), width=Pt(82))
para('Contents',size=20,bold=True,after=18)
p=para('')
r=p.add_run(); r._r.append(el('fldChar',fldCharType='begin',dirty='true'))
r=p.add_run(); instr=OxmlElement('w:instrText'); instr.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); instr.text=' TOC \\o "1-2" \\h \\z '; r._r.append(instr)
r=p.add_run(); r._r.append(el('fldChar',fldCharType='separate'))
p.add_run('Contents will be updated in Word.')
r=p.add_run(); r._r.append(el('fldChar',fldCharType='end'))
section_end({'top':51.85,'bottom':51.85,'left':57.6,'right':57.6})

buf=BytesIO(); d.save(buf)
with zipfile.ZipFile(SRC) as z: base={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(buf) as z: frag={n:z.read(n) for n in z.namelist()}
tree=E.fromstring(base['word/document.xml']); body=tree.find('w:body',NS)
original_children=[E.tostring(n) for n in body]
original_suffix=E.tostring(body)
front=E.fromstring(frag['word/document.xml']).find('w:body',NS)
rels=E.fromstring(base['word/_rels/document.xml.rels'])
frel=E.fromstring(frag['word/_rels/document.xml.rels'])
ids={r.get('Id') for r in rels}; mapping={}; additions={}
for rr in frel:
    if rr.get('Type').endswith('/image'):
        old=rr.get('Id'); target=rr.get('Target')
        rid='rIdNYXFront'+str(len(mapping)+1)
        assert rid not in ids
        name='nyx_front_'+Path(target).name
        mapping[old]=rid
        rr2=E.Element('{'+REL+'}Relationship',Id=rid,Type=rr.get('Type'),Target='media/'+name)
        rels.append(rr2); additions['word/media/'+name]=frag['word/'+target]

docid=max([int(x.get('id','0')) for x in tree.findall('.//wp:docPr',NS)]+[0])+1
prefix=[]
for n in list(front)[:-1]:
    n=deepcopy(n)
    for x in n.iter():
        for a,v in list(x.attrib.items()):
            if a.startswith('{'+NS['r']+'}') and v in mapping: x.set(a,mapping[v])
        if x.tag=='{'+NS['wp']+'}docPr': x.set('id',str(docid)); docid+=1
    prefix.append(n)
for i,n in enumerate(prefix): body.insert(i,n)
assert all(E.tostring(a)==b for a,b in zip(list(body)[len(prefix):],original_children)), 'Body content changed during prepend.'

# Retain original page furniture; restart original article at page 1.
sp=body[-1]
assert sp.tag==qn('w:sectPr')
typ=sp.find(qn('w:type'))
if typ is None:
    typ=el('type',val='nextPage')
    i=next((i for i,n in enumerate(sp) if E.QName(n).localname not in ('headerReference','footerReference','footnotePr','endnotePr')),0)
    sp.insert(i,typ)
else: typ.set(qn('w:val'),'nextPage')
pg=sp.find(qn('w:pgNumType'))
if pg is None:
    pg=el('pgNumType',fmt='decimal',start=1)
    cols=sp.find(qn('w:cols')); sp.insert(list(sp).index(cols) if cols is not None else len(sp),pg)
else: pg.set(qn('w:start'),'1'); pg.set(qn('w:fmt'),'decimal')

styles=E.fromstring(base['word/styles.xml'])
existing_ids={n.get(qn('w:styleId')) for n in styles}
for level in (1,2):
    sid='TOC'+str(level)
    assert sid not in existing_ids, 'Existing TOC style requires preservation-aware edit.'
    st=el('style',type='paragraph',styleId=sid)
    st.append(el('name',val='toc '+str(level))); st.append(el('basedOn',val='Normal'))
    st.append(el('next',val=sid)); st.append(el('autoRedefine')); st.append(el('uiPriority',val=39))
    pp=el('pPr'); pp.append(el('keepNext',val=0)); pp.append(el('keepLines'))
    tabs=el('tabs'); tabs.append(el('tab',val='right',leader='dot',pos=9936)); pp.append(tabs)
    pp.append(el('spacing',before=80 if level==1 else 0,after=40,line=240,lineRule='auto'))
    pp.append(el('ind',left=0 if level==1 else 240))
    st.append(pp)
    rp=el('rPr'); rp.append(el('rFonts',ascii='Times New Roman',hAnsi='Times New Roman',cs='Times New Roman'))
    rp.append(el('b',val=1 if level==1 else 0)); rp.append(el('color',val='000000')); rp.append(el('sz',val=22)); rp.append(el('szCs',val=22)); st.append(rp)
    styles.append(st)

ctype=E.fromstring(base['[Content_Types].xml'])
exts={x.get('Extension') for x in ctype if E.QName(x).localname=='Default'}
for ext,mime in [('png','image/png'),('jpeg','image/jpeg'),('jpg','image/jpeg')]:
    if ext not in exts: ctype.append(E.Element('{'+CT+'}Default',Extension=ext,ContentType=mime))

def xml(x): return E.tostring(x,encoding='UTF-8',xml_declaration=True,standalone=True)
out=base.copy(); out.update(additions)
out['word/document.xml']=xml(tree); out['word/_rels/document.xml.rels']=xml(rels)
out['word/styles.xml']=xml(styles); out['[Content_Types].xml']=xml(ctype)
OUT.parent.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
    for n,data in out.items(): z.writestr(n,data)
changed=[n for n in base if base[n]!=out[n]]
audit={'source_sha256':EXPECTED,'prefix_elements':len(prefix),'changed_parts':changed,
       'added_parts':list(additions),'source_preserved':hashlib.sha256(SRC.read_bytes()).hexdigest()==EXPECTED,
       'original_body_preserved_before_section_restart':True,
       'original_math_objects':len(tree.findall('.//m:oMath',NS)),
       'original_media_preserved':all(base[n]==out[n] for n in base if n.startswith('word/media/'))}
(TMP/'authoring_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
print(json.dumps(audit,indent=2)); print(OUT)
