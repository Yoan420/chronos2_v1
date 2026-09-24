from pathlib import Path
import sys,re,json,shutil,hashlib
ROOT=Path(__file__).resolve().parents[2]
TASK=Path(__file__).resolve().parent
OUT=ROOT/'output/nyx_paper_20260916'
SRC=ROOT/'tmp/tensor_timesfm_audit/metrics'
sys.path.insert(0,str(ROOT/'tmp/cv_these_75968/deps_readable'))
from latex2mathml.converter import convert
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import pandas as pd
XSLT=etree.XSLT(etree.parse('C:/Program Files (x86)/Microsoft Office/Root/Office16/MML2OMML.XSL'))
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.header_distance=Inches(.3);sec.footer_distance=Inches(.32)
for st in doc.styles:
    if st.type==1:
        st.font.color.rgb=RGBColor(0,0,0)
        for x in st.element.xpath('.//w:pBdr'):x.getparent().remove(x)
    if hasattr(st,'font'): st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.10;normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for name,size,space in [('Title',19,10),('Heading 1',13,11),('Heading 2',11.8,9)]:
    s=doc.styles[name];s.font.name='Cambria';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(0 if name=='Title' else space);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
    s.font.all_caps=False
cap=doc.styles['Caption'];cap.font.name='Times New Roman';cap.font.size=Pt(9.5);cap.font.italic=False;cap.font.bold=False;cap.font.color.rgb=RGBColor(0,0,0);cap.paragraph_format.space_after=Pt(8);cap.paragraph_format.line_spacing=1.04
head=sec.header.paragraphs[0]
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.core_properties.title='NYX for hourly day ahead electricity price forecasting across four European markets'
doc.core_properties.author='Yoan Kesraoui';doc.core_properties.subject='Retrospective experimental research manuscript';doc.core_properties.keywords='electricity price forecasting, NYX, Chronos-2, residual learning'

def add_inline(p,text):
    # Preserve real hyperlinks in references; regular prose is deliberately plain.
    pat=r'\[([^\]]+)\]\((https?://[^)]+)\)'
    pos=0
    for m in re.finditer(pat,text):
        p.add_run(text[pos:m.start()]);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(m.group(2),RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=m.group(1);r.append(t);h.append(r);p._p.append(h);pos=m.end()
    p.add_run(text[pos:])

def math(p,latex):
    tag=re.search(r'\\tag\{([^}]+)\}',latex)
    latex=re.sub(r'\\tag\{[^}]+\}','',latex).strip()
    # Word's MathML import elides large mspace; separate formula groups explicitly.
    for i,part in enumerate(latex.split(r'\qquad')):
        if i:p.add_run('     ')
        omml=XSLT(etree.fromstring(convert(part.strip()).encode()))
        root=omml.getroot()
        if root.tag==qn('m:oMathPara'):
            for child in list(root):p._p.append(child)
        else:p._p.append(root)
    if tag:p.add_run('   ('+tag.group(1)+')').font.size=Pt(10)

tables={}
def tabledef(key,caption,headers,rows,widths,note=''):
    tables[key]=dict(caption=caption,headers=headers,rows=rows,widths=widths,note=note)
fmt=lambda x:f'{float(x):.3f}'
signed=lambda x:f'{float(x):+.3f}'.replace('-','−')
data=pd.read_csv(OUT/'tables/data_description.csv')
tabledef('data_description','Table 1. Descriptive statistics of the evaluated hourly price series',['Zone','Mean','SD','Minimum','Median','Maximum','Negative hours'],[[r.zone,*[f'{v:.2f}' for v in [r['mean'],r.sd,r['min'],r['median'],r['max']]],str(r.negative_hours)] for _,r in data.iterrows()],[.5,.8,.8,.9,.85,.9,1.15],'Prices are in EUR/MWh; sample standard deviation. Each country has 8,760 observations over 365 delivery days. Negative hours have an observed price below zero.')
tabledef('inputs','Table 2. Evaluated inputs and their declared information timing',['Variable group','Unit','Use and timing'],[
['Country day-ahead prices','EUR/MWh','Historical hourly target; context through the end of D−1. Publication vintages are not independently certified.'],
['Residual-load forecasts for FR, DE, BE, NL and ES','GW','Five forecast channels, in historical context and future covariates; reconstructed D−1 08:00 cutoff.'],
['French nuclear-generation forecast','GW','Sixth forecast channel for every target country; same reconstructed cutoff.'],
['Calendar variables','Unitless','Deterministic hourly and calendar information.'],
['Storm reporting composite','EUR/MWh','External comparator only; cached forecasts plus recomputed missing periods.']],[1.65,.8,4.25],'No realized future fundamentals or Storm forecasts are intended as NYX inputs. The as-of metadata reflect the implemented query protocol, not verified provider publication timestamps.')
tabledef('configuration','Table 3. Effective settings of the evaluated pipeline',['Stage','Setting','Value'],[
['Chronos-2','Foundation weights; context','Frozen pinned checkpoint; 2,048 physical hours'],
['Chronos-2','Output; task interaction','Nine deciles requested; q10/q50/q90 retained; cross_learning=False'],
['CatBoost residual','Loss; daily fitting window','MAE; previous 365 civil days; minimum 720 rows'],
['CatBoost residual','Iterations; depth; learning rate','700; 6; 0.03'],
['CatBoost residual','L2; random seed; time order','15; 42; has_time=True'],
['CatBoost residual','Shift bound','±40 EUR/MWh, applied equally to each quantile'],
['Governed filter','Replay; candidates','Previous 365 civil days; five candidates'],
['Governed filter','Selection history; weight grid','At most 60 days; minimum 14; weights 0 to 1 by 0.05'],
['Governed filter','Required improvement; raw cap','max(0.05 EUR/MWh, 0.5% of upstream MAE); ±20 EUR/MWh']],[1.05,2.05,3.6],'There is no LoRA or foundation-model fine-tuning. Appendix B specifies filter dynamics. Settings are reported as executed, without treating unused configuration entries as active hyperparameters.')
s=pd.read_csv(TASK/'verified_stage_metrics.csv');s=s[(s.period=='annual')&(s.support=='four_model_common')]
zones=['BE','DE','FR','NL','pooled']
def val(z,model,col):return s[(s.zone==z)&(s.model==model)].iloc[0][col]
tabledef('annual_stages','Table 4. Annual MAE across the retained NYX stages',['Zone','Chronos-2','After CatBoost','NYX','MAE reduction (%)'],[[z.upper() if z=='pooled' else z,fmt(val(z,'chronos','mae')),fmt(val(z,'residual','mae')),fmt(val(z,'nyx','mae')),f"{100*(1-val(z,'nyx','mae')/val(z,'chronos','mae')):.2f}"] for z in zones],[.7,1.35,1.45,1.25,1.55],'MAE in EUR/MWh. All three stages use the same 8,759 hours per country as Table 5. Reduction is relative to the Chronos-2 stage. Pooled n = 35,036.')
tabledef('annual_benchmark','Table 5. Annual NYX and Storm reporting-composite performance',['Zone','NYX MAE','Storm MAE','NYX RMSE','Storm RMSE','NYX bias','Storm bias'],[[z.capitalize() if z=='pooled' else z,*[fmt(val(z,mod,metric)) for metric in ['mae','rmse','bias'] for mod in ['nyx','storm']]] for z in zones],[.64,1.01,1.01,1.01,1.01,1.01,1.01],'All errors in EUR/MWh on 35,036 paired observations. Bias = forecast − observation. Storm includes recomputed fallback hours; Table 9 checks sensitivity to their exclusion.')
e=pd.read_csv(OUT/'tables/extreme_regimes.csv');erows=[]
for z in zones[:4]:
    for reg,label in [('negative','Negative'),('high','Above initial q99')]:
        r=e[(e.zone==z)&(e.regime==reg)].iloc[0];erows.append([z,label,str(int(r.n)),f'{r.mae:.2f}',f'{r.bias:.2f}',f'{r.coverage80_pct:.2f}'])
tabledef('extremes','Table 6. NYX errors within realized extreme-price regimes',['Zone','Realized regime','Hours','MAE','Bias','Coverage (%)'],erows,[.5,1.7,.7,1.05,1.05,1.5],'Errors in EUR/MWh. Negative: y < 0. High: y exceeds its country-specific initial-sample q99, held fixed thereafter. P10–P90 coverage is conditional on the realized regime and is not a stand-alone test of conditional calibration.')
c=pd.read_csv(OUT/'tables/nyx_interval_calibration.csv');c=c[c.grouping=='all']
tabledef('intervals','Table 7. Annual NYX interval evaluation',['Zone','Coverage (%)','Mean width','Mean IS80','Mean WIS80'],[[r.zone,f'{100*r.coverage80:.2f}',f'{r.width80:.2f}',f'{r.interval_score80:.2f}',f'{r.wis80_single_interval:.3f}'] for _,r in c.iterrows()],[.7,1.5,1.5,1.5,1.5],'Coverage target: 80%; n = 8,760 per country. Width and scores are in EUR/MWh. WIS80 includes one interval plus the median and is not a complete-distribution score.')
t=pd.read_csv(OUT/'tables/locked_test_metrics.csv');t=t[(t.label=='latest')&(t.zone=='pooled')]
b=pd.read_csv(OUT/'tables/locked_test_paired_bootstrap.csv');b=b[(b.label=='latest')&(b.zone=='pooled')]
sel=json.loads((SRC/'validation_selection_locked.json').read_text())
selbest=sel['best_per_family'] if 'best_per_family' in sel else sel.get('selected_by_family',{})
if not selbest:
    selbest={}
    for r in sel['all_candidates']:
        k=r['spec']['family']
        if k not in selbest or r['mae']<selbest[k]['mae']:selbest[k]=r
names={'nyx':'NYX (identity)','ewma_country':'EWMA country (28 d)','ewma_country_hour':'EWMA country × hour (28 d)','ridge_univariate':'Univariate ridge (α = 100)','ridge_multivariate':'Multivariate ridge (α = 100)','pca_ridge':'PCA ridge (rank 1, α = 100)','storm':'Storm reporting composite'}
rows=[]
for _,r in t.iterrows():
    bb=b[b.model==r.model].iloc[0];v='—' if r.model=='storm' else fmt(selbest[r.model]['mae']);ci='0 (reference)' if r.model=='nyx' else f"{signed(bb.mae_delta)}\n[{signed(bb.mae_delta_ci_low)}, {signed(bb.mae_delta_ci_high)}]"
    rows.append([names[r.model],v,fmt(r.mae),fmt(r.rmse),fmt(r.bias),ci])
tabledef('residual_test','Table 8. Validation-selected residual configurations on the reserved 90 days',['Model and selected setting','Validation MAE','Test MAE','Test RMSE','Test bias','ΔMAE [95% interval]'],rows,[1.7,.85,.82,.85,.78,1.6],'n = 8,640 country–hours in the reserved segment. Errors in EUR/MWh; positive ΔMAE means worse than NYX. Validation uses frozen labels; test uses refreshed labels. Paired moving-block intervals use 2,000 replicates and seven-day blocks. Storm is not part of residual-model selection.')
p=pd.read_csv(TASK/'storm_provenance_sensitivity.csv');p=p[p.zone=='pooled'];rows=[]
for period,label in [('annual','Annual'),('test90','Reserved 90 d')]:
    for subset,sl in [('all_common','All common hours'),('cache_only','Exclude Storm recomputation'),('cache_only_no_label_fallback','Also exclude EPEX label fallback')]:
        a=p[(p.period==period)&(p.subset==subset)&(p.model=='nyx')].iloc[0];b1=p[(p.period==period)&(p.subset==subset)&(p.model=='storm')].iloc[0]
        rows.append([label,sl,f'{int(a.n):,}',fmt(a.mae),fmt(b1.mae)])
tabledef('provenance_sensitivity','Table 9. Pooled MAE sensitivity to source provenance',['Period','Included support','Pairs','NYX MAE','Storm MAE'],rows,[1.1,2.3,1,1.15,1.05],'Errors in EUR/MWh. Samples are matched within each row and weighted by country–hour. Excluding fallback values does not independently certify the remaining source vintages.')
tabledef('filters','Table 10. Candidate additive corrections relative to the upstream median',['Candidate','State size','Proposed correction'],[
['Linear bias','1','MATH:b_0'],['Linear harmonic','3',r'MATH:b_0+a_s\sin(2\pi h/24)+a_c\cos(2\pi h/24)'],['Linear market','14','Linear combination of intercept, hour sine/cosine and 11 standardized market inputs'],['Linear scale','2',r'MATH:b_0+(s-1)a,\quad 0.5\leq s\leq1.5'],['Unscented scale','2',r'MATH:b_0+(s(u)-1)a,\quad s(u)=0.5+(1+e^{-u})^{-1}']],[1.2,.8,4.7],'The 11 market inputs are upstream median, interval width, CatBoost shift, five residual-load forecasts, their mean and range, and the French nuclear forecast. Symbols in this table describe candidate states; the final bounded weighted shift is b in Equation (11).')

def put_table(key):
    spec=tables[key]
    p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;p.add_run(spec['caption']).bold=True
    tb=doc.add_table(rows=1,cols=len(spec['headers']));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    # Explicit table geometry and borders, independent of Word theme defaults.
    pr=tb._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    pr.append(borders)
    for col,w in zip(tb.columns,spec['widths']):col.width=Inches(w)
    for j,h in enumerate(spec['headers']):tb.rows[0].cells[j].text=h
    repeat=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in spec['rows']:
        cells=tb.add_row().cells
        for j,value in enumerate(row):
            p1=cells[j].paragraphs[0]
            if str(value).startswith('MATH:'):math(p1,str(value)[5:])
            else:p1.add_run(str(value))
    for i,row in enumerate(tb.rows):
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        for j,cell in enumerate(row.cells):
            cell.width=Inches(spec['widths'][j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side,value in [('top','65'),('bottom','65'),('left','80'),('right','80')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');marg.append(el)
            cp.append(marg)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EBEE');cp.append(sh)
            for p1 in cell.paragraphs:
                p1.paragraph_format.space_after=Pt(0);p1.paragraph_format.space_before=Pt(0);p1.paragraph_format.line_spacing=1.02
                p1.paragraph_format.keep_with_next=(i==0 or i==len(tb.rows)-1 or key in ['residual_test','configuration'])
                prose=key in ['inputs','configuration','filters']
                p1.alignment=WD_ALIGN_PARAGRAPH.LEFT if (j==0 or prose or (key=='provenance_sensitivity' and j<2)) else WD_ALIGN_PARAGRAPH.CENTER
                for run in p1.runs:run.font.size=Pt(9.4);run.font.name='Times New Roman';run.font.bold=(i==0);run.font.color.rgb=RGBColor(0,0,0)
    if spec['note']:
        p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=1.02
        p.add_run(spec['note']).font.size=Pt(9)

source=(OUT/'NYX_Research_Manuscript.md').read_text(encoding='utf-8')
in_refs=False
for block in re.split(r'\n\s*\n',source.strip()):
    block=block.strip()
    if block.startswith('# '):doc.add_paragraph(block[2:],style='Title')
    elif block.startswith('## '):
        heading=block[3:];p=doc.add_paragraph(heading,style='Heading 1')
        if heading in ['1 Introduction','References','Appendix A Reproducibility details']:p.paragraph_format.page_break_before=True
        in_refs=heading=='References'
    elif block.startswith('### '):doc.add_paragraph(block[4:],style='Heading 2')
    elif block.startswith('@table:'):put_table(block.split(':',1)[1])
    elif block.startswith('@figure:'):
        name,caption=block[8:].split('|',1);p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(OUT/'figures'/(name+'.png')),width=Inches(6.7))
        p=doc.add_paragraph(caption,style='Caption')
    elif block.startswith('$$'):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(9);p.paragraph_format.space_before=Pt(3);math(p,block[2:-2])
    else:
        p=doc.add_paragraph();add_inline(p,block.replace('\n',' '))
        if in_refs:p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22);p.paragraph_format.line_spacing=1.05
        elif block.startswith(('Yoan Kesraoui','Correspondence:','Research manuscript')):
            p.paragraph_format.space_after=Pt(4)
            for run in p.runs:run.font.size=Pt(10.5);run.bold=block=='Yoan Kesraoui'
        elif block.startswith('Keywords:'):
            for run in p.runs:run.font.size=Pt(10)
        else:p.alignment=WD_ALIGN_PARAGRAPH.LEFT if block.startswith('The retained foundation checkpoint') else WD_ALIGN_PARAGRAPH.JUSTIFY
for p in doc.paragraphs:
    for el in p._p.xpath('.//w:pBdr'):el.getparent().remove(el)
doc.save(OUT/'NYX_Research_Manuscript.docx')
(OUT/'tables/manuscript_tables.json').write_text(json.dumps(tables,indent=2,ensure_ascii=False),encoding='utf-8')
for name in ['verified_stage_metrics.csv','stage_paired_block_bootstrap.csv','weekly_naive_matched_metrics.csv','weekly_naive_paired_bootstrap.csv','storm_provenance_sensitivity.csv','source_fallback_counts.csv','verified_interval_metrics.csv']:
    shutil.copyfile(TASK/name,OUT/'tables'/name)
print('DOCX created; words',len(source.split()),'tables',len(tables),'figures',source.count('@figure:'))
